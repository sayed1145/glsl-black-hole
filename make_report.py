"""Build a modern DOCX technical report from the included Markdown and figures.
Usage: python make_report.py --root . --kind meadow
"""
from pathlib import Path
import re,argparse
from docx import Document
from docx.shared import Inches,Pt,RGBColor,Cm
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def clean(s):
    return re.sub(r'\[([^\]]+)\]\(([^)]+)\)',r'\1 (\2)',s).replace('**','').replace('`','')

def shade(cell,color):
    p=cell._tc.get_or_add_tcPr();e=OxmlElement('w:shd');e.set(qn('w:fill'),color);p.append(e)

def build(root,kind):
    root=Path(root).resolve();folder=root/'report';folder.mkdir(exist_ok=True)
    meadow=kind=='meadow';accent='335B35' if meadow else '856D4C'
    title='程序化草地与山谷' if meadow else '程序化黑洞'
    filename='草地_实现与验证报告.docx' if meadow else '黑洞_实现与验证报告.docx'
    doc=Document();sec=doc.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.0);sec.right_margin=Cm(2.0)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Aptos';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        if name=='Normal':style.font.size=Pt(10.5);style.paragraph_format.space_after=Pt(6);style.paragraph_format.line_spacing=1.18
        if name.startswith('Heading'):style.font.color.rgb=RGBColor.from_string(accent)
    header=sec.header.paragraphs[0];header.text='GLSL STUDY  /  '+title;header.style=doc.styles['Caption']
    header.runs[0].font.color.rgb=RGBColor.from_string(accent)
    footer=sec.footer.paragraphs[0];footer.alignment=2
    footer.add_run('技术报告  ·  ')
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    p=doc.add_paragraph('PROCEDURAL RENDERING / 纯着色器重建');p.runs[0].font.size=Pt(10);p.runs[0].font.color.rgb=RGBColor.from_string(accent)
    doc.add_heading(title,0);doc.add_paragraph('实现、对照、验证与复现',style='Subtitle')
    poster=folder/'rendered.jpg'
    if poster.exists():doc.add_picture(str(poster),width=Inches(6.25 if meadow else 5.4))
    summary='960 × 600  /  8 秒  /  20 fps  /  160 帧' if meadow else '960 × 960  /  8 秒  /  25 fps  /  200 帧'
    p=doc.add_paragraph(summary);p.runs[0].bold=True
    doc.add_paragraph('视觉近似重建，不是逐像素复制。参考图仅用于离线校准与本报告对照，运行时不作为纹理输入。')
    doc.add_page_break()
    doc.add_heading('参考图与重建结果',1)
    table=doc.add_table(rows=2,cols=2);table.autofit=False
    for i,label in enumerate(['用户参考图','GLSL 重建首帧']):
        table.cell(0,i).text=label;shade(table.cell(0,i),'EDF2EA' if meadow else 'F2EEE8')
    for i,path in enumerate([folder/'reference.jpg',folder/'rendered.jpg']):
        if path.exists():table.cell(1,i).paragraphs[0].add_run().add_picture(str(path),width=Inches(3.0))
    doc.add_paragraph('对照图不表示已还原原始场景。构图、主色调和效果类型作为拟合目标；局部纹理、几何与噪声细节仍有差异。')
    markdown=(folder/'报告.md').read_text()
    lines=markdown.splitlines();i=0;code=False
    while i<len(lines):
        s=lines[i].strip()
        if s.startswith('```'):code=not code;i+=1;continue
        if not s:i+=1;continue
        if s.startswith('|') and i+1<len(lines) and re.match(r'^\|[\s:\-|]+\|$',lines[i+1].strip()):
            rows=[[clean(x.strip()) for x in s.strip('|').split('|')]];i+=2
            while i<len(lines) and lines[i].strip().startswith('|'):
                rows.append([clean(x.strip()) for x in lines[i].strip().strip('|').split('|')]);i+=1
            t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Table Grid'
            for ri,row in enumerate(rows):
                cells=t.add_row().cells
                for ci,text in enumerate(row[:len(cells)]):
                    cells[ci].text=text
                    if ri==0:
                        shade(cells[ci],accent)
                        for run in cells[ci].paragraphs[0].runs:run.font.color.rgb=RGBColor(255,255,255);run.bold=True
                    elif ri%2==0:shade(cells[ci],'F3F5F2')
            doc.add_paragraph();continue
        if code:
            p=doc.add_paragraph(s);p.paragraph_format.space_after=Pt(1)
            for run in p.runs:run.font.name='Consolas';run.font.size=Pt(8.5)
        elif s.startswith('# '):pass
        elif s.startswith('### '):doc.add_heading(clean(s[4:]),2)
        elif s.startswith('## '):doc.add_heading(clean(s[3:]),1)
        elif s.startswith('- '):doc.add_paragraph(clean(s[2:]),style='List Bullet')
        else:doc.add_paragraph(clean(s))
        i+=1
    if meadow:
        doc.add_heading('附图：效果开关检查',1)
        doc.add_paragraph('下图为同一相机的诊断输出。几何阴影关闭时仍保留云影和根部 AO；泛光关闭时仍保留大气雾。')
        t=doc.add_table(rows=2,cols=2)
        for i,(label,name) in enumerate([('几何投影阴影关闭','shadows_off.jpg'),('泛光关闭','bloom_off.jpg')]):
            t.cell(0,i).text=label
            if (folder/name).exists():t.cell(1,i).paragraphs[0].add_run().add_picture(str(folder/name),width=Inches(3.0))
    path=folder/filename;doc.save(path);print(path)
    return path

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--root',default=str(Path(__file__).resolve().parent));ap.add_argument('--kind',choices=['meadow','black_hole'],default='meadow');a=ap.parse_args();build(a.root,a.kind)
